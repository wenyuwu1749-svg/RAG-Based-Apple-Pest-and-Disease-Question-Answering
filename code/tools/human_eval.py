"""Blinded human scoring: export a stratified sample, import filled sheets, compute agreement.

python code/tools/human_eval.py export --name test_v9 --bank data/eval/test_questions.xlsx --n 60 --raters "Wenyu Wu" "Dongyang Chen" "Mengshuai Feng"
python code/tools/human_eval.py import --name test_v9
python code/tools/human_eval.py agreement --name test_v9

export: stratified random sample over model x condition (equal allocation, at most one answer per
  question per stratum), only repeat 1. The rater sheet hides model, condition, repeat and judge
  score, removes citation markers like [1], and shuffles rows independently for each rater.
  The answer key (item -> task) is written to a separate file that raters must not open.
import: reads every filled rater sheet; rows without a 0-4 score are reported, never filled in.
agreement: exact agreement, within-one agreement, quadratic-weighted Cohen's kappa and confusion
  matrices for every rater pair and for each rater vs the automatic judge, plus mean differences.
Nothing in this script ever writes a human score.
"""
import argparse
import csv
import json
import random
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'code' / 'tools'))
import question_bank as qb  # noqa: E402

LOGS = ROOT / 'data' / 'logs'
OUT = ROOT / 'data' / 'eval' / 'human_eval'
RUBRIC = [('4', 'Core answer (diagnosis or main fact) correct; advice complete and correct'),
          ('3', 'Core answer correct; main advice correct, only secondary details missing'),
          ('2', 'Core answer correct, but advice sparse or generic'),
          ('1', 'Core answer uncertain or only partly correct, but some useful correct information'),
          ('0', 'Wrong core answer, any listed major error, or no useful content')]


def strip_citations(text):
    return re.sub(r'\s*\[(\d+(\s*[,\-–]\s*\d+)*)\]', '', text or '')


def safe(name):
    return re.sub(r'[^A-Za-z0-9]+', '', name)


def cmd_export(a):
    bank = {r['question_id']: r for r in qb.read_rows(ROOT / a.bank)}
    answers = [x for x in csv.DictReader((LOGS / f'{a.name}_answers.csv').open(encoding='utf-8-sig'))
               if x['status'] == 'success' and str(x['repeat']) == '1']
    strata = defaultdict(list)
    for x in answers:
        strata[(x['model'], x['condition'])].append(x)
    rng = random.Random(a.seed)
    per = a.n // len(strata)
    sample = []
    for key in sorted(strata):
        pool = strata[key][:]
        rng.shuffle(pool)
        sample += pool[:per]
    rng.shuffle(sample)
    OUT.mkdir(parents=True, exist_ok=True)
    key_rows = [{'item_id': f'H{i:03d}', 'task_id': x['task_id'], 'model': x['model'], 'condition': x['condition'],
                 'question_id': x['question_id']} for i, x in enumerate(sample, 1)]
    with (OUT / f'{a.name}_ANSWER_KEY_do_not_share.csv').open('w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=list(key_rows[0])); w.writeheader(); w.writerows(key_rows)
    for rater in a.raters:
        order = list(zip(key_rows, sample))
        random.Random(f'{a.seed}-{rater}').shuffle(order)
        wb = Workbook(); ws = wb.active; ws.title = 'scoring'
        head = ['item_id', 'question', 'reference_answer', 'required_points', 'major_errors', 'answer_to_score', 'human_score_0_4', 'comment']
        widths = [8, 40, 55, 40, 30, 70, 12, 30]
        for c, (h, wdt) in enumerate(zip(head, widths), 1):
            cell = ws.cell(1, c, h); cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor='FFF2CC' if c >= 7 else 'DDEBF7')
            ws.column_dimensions[cell.column_letter].width = wdt
        dv = DataValidation(type='whole', operator='between', formula1='0', formula2='4', allow_blank=True)
        ws.add_data_validation(dv)
        for r, (k, x) in enumerate(order, 2):
            q = bank[x['question_id']]
            vals = [k['item_id'], q['question_en'], q['reference_answer_en'],
                    '\n'.join('• ' + p for p in qb.split_list(q['required_points'])),
                    '\n'.join('• ' + p for p in qb.split_list(q['major_errors'])), strip_citations(x['answer']), None, None]
            for c, v in enumerate(vals, 1):
                ws.cell(r, c, v).alignment = Alignment(wrap_text=True, vertical='top')
            dv.add(f'G{r}')
        ws.freeze_panes = 'B2'
        g = wb.create_sheet('instructions')
        lines = [f'Rater: {rater}', '', 'Score every row 0-4 in column G using only this rubric:'] + [f'{s} = {t}' for s, t in RUBRIC] + [
            '', 'Rules:', '- Work alone; do not discuss items with other raters until all sheets are returned.',
            '- The model, condition and automatic score are hidden on purpose; do not try to guess them.',
            '- Judge content against the reference answer and required points. Correct advice that differs from the reference but serves the same purpose counts.',
            '- Overseas spray timing or products presented as New Zealand advice count as a major error when the question is about New Zealand.',
            '- Citations and source names are neither rewarded nor penalised.',
            '- Leave a cell empty if you cannot decide, and explain in the comment column.',
            '- Record roughly how many minutes the whole sheet took in the first comment cell.']
        for i, t in enumerate(lines, 1):
            g.cell(i, 1, t)
        g.column_dimensions['A'].width = 130
        path = OUT / f'{a.name}_rater_{safe(rater)}.xlsx'
        wb.save(path)
        print('wrote', path.relative_to(ROOT))
    print(f'{len(sample)} items from {len(strata)} strata ({per} each); answer key kept separately.')


def read_scores(name):
    key = {r['item_id']: r for r in csv.DictReader((OUT / f'{name}_ANSWER_KEY_do_not_share.csv').open(encoding='utf-8-sig'))}
    scores, problems = {}, []
    for path in sorted(OUT.glob(f'{name}_rater_*.xlsx')):
        rater = path.stem.split('_rater_')[-1]
        ws = load_workbook(path, data_only=True)['scoring']
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            v = row[6]
            if v is None or str(v).strip() == '':
                problems.append(f'{rater}: {row[0]} has no score')
                continue
            try:
                s = int(float(v))
                assert 0 <= s <= 4
            except Exception:
                problems.append(f'{rater}: {row[0]} invalid score {v!r}')
                continue
            scores.setdefault(rater, {})[row[0]] = s
    return key, scores, problems


def cmd_import(a):
    key, scores, problems = read_scores(a.name)
    for p in problems:
        print('WARNING', p)
    judged = {}
    jf = LOGS / f'{a.name}_judged.jsonl'
    if jf.exists():
        for line in jf.read_text(encoding='utf-8').splitlines():
            j = json.loads(line)
            if 'score' in j:
                judged[j['task_id']] = j['score']
    out = OUT / f'{a.name}_human_scores.csv'
    raters = sorted(scores)
    with out.open('w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(['item_id', 'task_id', 'model', 'condition', 'question_id', 'judge_score'] + raters)
        for item, k in sorted(key.items()):
            w.writerow([item, k['task_id'], k['model'], k['condition'], k['question_id'], judged.get(k['task_id'], '')] +
                       [scores[r].get(item, '') for r in raters])
    print(f'raters: {raters}; complete items per rater: {[len(scores[r]) for r in raters]} of {len(key)}. Saved {out.relative_to(ROOT)}')


def weighted_kappa(x, y, k=5):
    n = len(x)
    if n == 0:
        return None
    obs = [[0] * k for _ in range(k)]
    for a, b in zip(x, y):
        obs[a][b] += 1
    ra = [sum(obs[i]) for i in range(k)]
    cb = [sum(obs[i][j] for i in range(k)) for j in range(k)]
    num = sum(((i - j) ** 2) * obs[i][j] for i in range(k) for j in range(k))
    den = sum(((i - j) ** 2) * ra[i] * cb[j] / n for i in range(k) for j in range(k))
    return 1 - num / den if den else None


def compare(label, x, y):
    n = len(x)
    exact = sum(a == b for a, b in zip(x, y)) / n
    within = sum(abs(a - b) <= 1 for a, b in zip(x, y)) / n
    diff = sum(b - a for a, b in zip(x, y)) / n
    kappa = weighted_kappa(x, y)
    cm = [[sum(1 for a, b in zip(x, y) if a == i and b == j) for j in range(5)] for i in range(5)]
    return {'pair': label, 'n': n, 'exact_agreement': round(exact, 3), 'within_one': round(within, 3),
            'quadratic_weighted_kappa': round(kappa, 3) if kappa is not None else '', 'mean_difference_second_minus_first': round(diff, 3),
            'confusion_matrix_rows_first_cols_second': json.dumps(cm)}


def cmd_agreement(a):
    path = OUT / f'{a.name}_human_scores.csv'
    rows = list(csv.DictReader(path.open(encoding='utf-8-sig')))
    raters = [c for c in rows[0] if c not in ('item_id', 'task_id', 'model', 'condition', 'question_id', 'judge_score')]
    results = []
    sources = raters + ['judge_score']
    for i, r1 in enumerate(sources):
        for r2 in sources[i + 1:]:
            pairs = [(int(x[r1]), int(x[r2])) for x in rows if x[r1] != '' and x[r2] != '']
            if pairs:
                results.append(compare(f'{r1} vs {r2}', [p[0] for p in pairs], [p[1] for p in pairs]))
    hm = [x for x in rows if all(x[r] != '' for r in raters) and x['judge_score'] != '']
    if hm and raters:
        mean_h = [round(sum(int(x[r]) for r in raters) / len(raters)) for x in hm]
        results.append(compare('human mean (rounded) vs judge_score', mean_h, [int(x['judge_score']) for x in hm]))
    out = OUT / f'{a.name}_agreement.csv'
    with out.open('w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=list(results[0])); w.writeheader(); w.writerows(results)
    for r in results:
        print(f"{r['pair']:<45} n={r['n']:<3} exact {r['exact_agreement']}  within1 {r['within_one']}  kappa_w {r['quadratic_weighted_kappa']}  diff {r['mean_difference_second_minus_first']}")
    print('Saved', out.relative_to(ROOT))


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest='cmd', required=True)
    e = sub.add_parser('export')
    e.add_argument('--name', required=True); e.add_argument('--bank', required=True)
    e.add_argument('--n', type=int, default=60); e.add_argument('--seed', type=int, default=20261021)
    e.add_argument('--raters', nargs='+', required=True)
    for c in ('import', 'agreement'):
        sub.add_parser(c).add_argument('--name', required=True)
    a = ap.parse_args()
    {'export': cmd_export, 'import': cmd_import, 'agreement': cmd_agreement}[a.cmd](a)


if __name__ == '__main__':
    main()
