"""Export the development and test question banks into one workbook (sheet 1 notes, sheet 2 dev, sheet 3 test).

python code/tools/export_question_bank.py
Reads data/eval/questions/dev_questions.xlsx and test_questions.xlsx (unchanged) and writes data/eval/questions/G6_question_bank.xlsx.
"""
import collections
import hashlib
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'code' / 'tools'))
import question_bank as qb  # noqa: E402

OUT = ROOT / 'data' / 'eval' / 'questions' / 'G6_question_bank.xlsx'
FIELDS = [('question_id', 'Question ID', 9), ('question_type', 'Question type', 18), ('answerability', 'Answerability', 12), ('style', 'Question style', 12),
          ('topic_id', 'Topic', 7), ('region', 'Region', 18), ('stage', 'Stage', 12), ('question_en', 'Question', 50),
          ('reference_answer_en', 'Reference answer', 70), ('required_points', 'Required points', 50), ('major_errors', 'Major errors', 40),
          ('acceptable_partial', 'Answerable / to decline', 35), ('source_ids', 'Source IDs', 18), ('chunk_ids', 'Cited chunks', 30),
          ('location', 'Location in source', 30), ('reviewer', 'Reviewer', 14), ('status', 'Status', 9)]
STYLE = {"Farmer's voice": "Farmer's voice", "With region": "With region or environment context", "Rare topic": "Rare topic"}


def style_of(r):
    note = str(r.get('notes') or '')
    return next((v for k, v in STYLE.items() if note.startswith(k)), '')


def sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    head_fill = PatternFill('solid', fgColor='DDEBF7')
    for c, (_, h, w) in enumerate(FIELDS, 1):
        cell = ws.cell(1, c, h)
        cell.font, cell.fill = Font(bold=True), head_fill
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[cell.column_letter].width = w
    for i, r in enumerate(rows, 2):
        for c, (key, _, _) in enumerate(FIELDS, 1):
            v = style_of(r) if key == 'style' else r.get(key, '')
            if key in ('required_points', 'major_errors'):
                v = '\n'.join('• ' + p for p in qb.split_list(v))
            elif key in ('source_ids', 'chunk_ids'):
                v = '\n'.join(qb.split_list(v))
            ws.cell(i, c, v).alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = ws.dimensions


def main():
    dev = qb.read_rows(ROOT / 'data' / 'eval' / 'questions' / 'dev_questions.xlsx')
    test = qb.read_rows(ROOT / 'data' / 'eval' / 'questions' / 'test_questions.xlsx')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Notes'
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 100

    def count(rows, f):
        return '; '.join(f'{k} {v}' for k, v in collections.Counter(f(r) for r in rows).most_common())

    lines = [
        ('G6 apple pest and disease RAG question bank', None),
        ('', None),
        ('Purpose', 'Evaluate the accuracy of general-purpose LLMs answering apple pest and disease questions under no RAG, standard RAG and token-efficient RAG (plus the headers-only ablation).'),
        ('Sheets', '"Development set" (20 questions) is used to tune the system and choose the judge and parameters; "Test set" (50 questions) is used only to answer the research questions (the preliminary and formal experiments use the same test questions).'),
        ('Question style', 'Symptoms described in a farmer\'s voice without naming the cause and without unnecessary conditions; the reference answer gives the most likely cause, look-alike problems and how to tell them apart, and the management advice supported by the sources.'),
        ('Development set', f'{len(dev)} questions; answerability: {count(dev, lambda r: r["answerability"])}; style: {count(dev, style_of)}'),
        ('Test set', f'{len(test)} questions; answerability: {count(test, lambda r: r["answerability"])}; style: {count(test, style_of)}'),
        ('Test question types', count(test, lambda r: r['question_type'])),
        ('', None),
        ('Fields', None),
        ('Question ID', 'D = development set, T = test set'),
        ('Question type', 'symptoms_identification; management_strategy; conditions_risk; monitoring_sampling; postharvest; boundary_insufficient (evidence insufficient)'),
        ('Answerability', 'answerable; partial; unanswerable from the corpus (the correct answer states the limits and does not invent)'),
        ('Topic', 'Scope table identifiers T01–T46'),
        ('Reference answer', 'Answer written only from the cited source passages'),
        ('Required points', '3–5 points a complete answer must contain; basis for scoring'),
        ('Major errors', 'Errors that score 0, e.g. misdiagnosis or presenting overseas spray advice as New Zealand advice'),
        ('Answerable / to decline', 'For partly answerable and unanswerable questions: what can be answered, what should be declined as unsupported, and why'),
        ('Source IDs / cited chunks', 'Source and chunk identifiers in the frozen corpus corpus-v1-20261014; traceable to the original text in the database'),
        ('Reviewer / status', 'Reviewed by the three team members; test set status frozen, development set reviewed'),
        ('', None),
        ('Data source', 'Exported by code/tools/export_question_bank.py from data/eval/questions/dev_questions.xlsx and test_questions.xlsx without changing the content.'),
        ('Source file SHA-256', f"dev {hashlib.sha256((ROOT / 'data/eval/questions/dev_questions.xlsx').read_bytes()).hexdigest()[:16]}…; test {hashlib.sha256((ROOT / 'data/eval/questions/test_questions.xlsx').read_bytes()).hexdigest()[:16]}…"),
        ('Change log', 'See data/eval/question_review/*/'),
    ]
    for i, (k, v) in enumerate(lines, 1):
        a = ws.cell(i, 1, k)
        if v is None:
            a.font = Font(bold=True, size=14 if i == 1 else 11)
        else:
            a.font = Font(bold=True)
            ws.cell(i, 2, v).alignment = Alignment(wrap_text=True, vertical='top')
    sheet(wb, 'Development set', dev)
    sheet(wb, 'Test set', test)
    wb.save(OUT)
    print('wrote', OUT.relative_to(ROOT), f'(dev {len(dev)}, test {len(test)})')


if __name__ == '__main__':
    main()
